"""Excel export. With an .xlsx template, cases are written into a *copy of that workbook*: header, styling,
column widths, dropdowns and other sheets are kept; example rows below the header are replaced.

Every value is written as text (never as a formula): crawled text such as "=HYPERLINK(...)" must not
execute when the file is opened (spreadsheet formula injection).
"""

from __future__ import annotations

from collections.abc import Sequence
from copy import copy
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from scoutqa.export.report import COVERAGE_HEADERS, LIMITATION_HEADERS, TRACE_HEADERS, Report
from scoutqa.generate.cases import TestCase
from scoutqa.template.spec import CanonicalField, TemplateSpec
from scoutqa.template.values import rows_for

MAX_CELL = 32_000  # Excel's limit is 32,767 characters
_HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_THIN = Side(style="thin", color="BFBFBF")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_WRAP = Alignment(wrap_text=True, vertical="top")
_WIDTHS = {
    CanonicalField.ID: 16, CanonicalField.MODULE: 16, CanonicalField.TITLE: 48, CanonicalField.SCENARIO: 36,
    CanonicalField.PRECONDITIONS: 32, CanonicalField.STEPS: 64, CanonicalField.EXPECTED: 48,
    CanonicalField.PRIORITY: 11, CanonicalField.TYPE: 13, CanonicalField.TEST_DATA: 36, CanonicalField.NOTES: 40,
    CanonicalField.STEP_NO: 8, CanonicalField.STEP_ACTION: 56, CanonicalField.STEP_EXPECTED: 44,
}


def _set(cell: Cell | MergedCell, value: str | int) -> None:
    if isinstance(value, str):
        cell.value = value[:MAX_CELL]
        cell.data_type = "s"  # literal text, even if it starts with '='
    else:
        cell.value = value


def _style_header(ws: Worksheet, row: int, count: int) -> None:
    for col in range(1, count + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill, cell.font, cell.border = _HEADER_FILL, _HEADER_FONT, _BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def _table_sheet(wb: Workbook, title: str, headers: list[str], rows: Sequence[Sequence[str | int]],
                 widths: list[int]) -> None:
    name = title if title not in wb.sheetnames else f"ScoutQA {title}"
    ws = wb.create_sheet(name)
    for col, header in enumerate(headers, 1):
        _set(ws.cell(row=1, column=col), header)
    _style_header(ws, 1, len(headers))
    for r, row in enumerate(rows, 2):
        for col, value in enumerate(row, 1):
            cell = ws.cell(row=r, column=col)
            _set(cell, value)
            cell.alignment, cell.border = _WRAP, _BORDER
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"


def export_excel(cases: list[TestCase], template: TemplateSpec, report: Report, out: Path) -> int:
    """Write the workbook; returns the number of data rows written to the test case sheet."""
    rows = [row for case in cases for row in rows_for(case, template)]
    if template.format == "xlsx":
        wb = load_workbook(template.source)
        sheet = wb[template.sheet] if template.sheet else wb.worksheets[0]
        assert isinstance(sheet, Worksheet), "the template sheet must be a worksheet"
        ws = sheet
        header_row = template.header_row
        styles = _data_row_styles(ws, header_row, len(template.columns))
        if ws.max_row > header_row:
            ws.delete_rows(header_row + 1, ws.max_row - header_row)
    else:
        wb = Workbook()
        active = wb.active
        assert isinstance(active, Worksheet)
        ws = active
        ws.title = "Test Cases"
        header_row, styles = 1, {}
        for col_spec in template.columns:
            _set(ws.cell(row=1, column=col_spec.index + 1), col_spec.name)
        _style_header(ws, 1, max((c.index for c in template.columns), default=0) + 1)
        for col_spec in template.columns:
            letter = get_column_letter(col_spec.index + 1)
            ws.column_dimensions[letter].width = _WIDTHS.get(col_spec.field, 22)
            if col_spec.enum:
                dv = DataValidation(type="list", formula1='"' + ",".join(col_spec.enum) + '"', allow_blank=True)
                dv.add(f"{letter}2:{letter}{max(len(rows) + 1, 2)}")
                ws.add_data_validation(dv)
        ws.freeze_panes = "B2"

    first = header_row + 1
    for r, row in enumerate(rows, first):
        for col_spec, value in zip(template.columns, row, strict=True):
            cell = ws.cell(row=r, column=col_spec.index + 1)
            _set(cell, value)
            if col_spec.index in styles:
                _apply(cell, styles[col_spec.index])
            else:
                cell.alignment, cell.border = _WRAP, _BORDER
    last = first + len(rows) - 1
    if template.format == "xlsx":
        _extend_validations(ws, header_row, last)
    if rows and not ws.auto_filter.ref:
        width = max(c.index for c in template.columns) + 1
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(width)}{max(last, header_row)}"

    _table_sheet(wb, "Coverage", COVERAGE_HEADERS, report.coverage, [16, 28, 28, 16, 10, 10, 8, 8, 10, 8, 10, 10])
    _table_sheet(wb, "Limitations", LIMITATION_HEADERS, report.limitations, [30, 70, 8, 60])
    _table_sheet(wb, "Trace", TRACE_HEADERS, report.trace, [16, 50, 8, 22, 24, 14, 18, 50, 50, 10])
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return len(rows)


# ---------------------------------------------------------------- template-preserving helpers

_Style = tuple[Any, Any, Any, Any, str]  # font, fill, border, alignment, number format


def _data_row_styles(ws: Worksheet, header_row: int, count: int) -> dict[int, _Style]:
    """Style of the template's first data row (if the template has one) per 0-based column."""
    styles: dict[int, _Style] = {}
    row = header_row + 1
    if row > ws.max_row:
        return styles
    for col in range(1, max(count, ws.max_column) + 1):
        cell = ws.cell(row=row, column=col)
        if cell.has_style:
            alignment = copy(cell.alignment)
            alignment.wrap_text = True
            styles[col - 1] = (copy(cell.font), copy(cell.fill), copy(cell.border), alignment, cell.number_format)
    return styles


def _apply(cell: Cell | MergedCell, style: _Style) -> None:
    cell.font, cell.fill, cell.border, cell.alignment, cell.number_format = style


def _extend_validations(ws: Worksheet, header_row: int, last_row: int) -> None:
    """Stretch the template's dropdowns (data validations) over all written rows."""
    for dv in ws.data_validations.dataValidation:
        for rng in list(dv.sqref.ranges):
            if rng.min_row is not None and rng.min_row > header_row and last_row > (rng.max_row or 0):
                letter_min, letter_max = get_column_letter(rng.min_col), get_column_letter(rng.max_col)
                dv.sqref.remove(rng)
                dv.sqref.add(f"{letter_min}{rng.min_row}:{letter_max}{last_row}")
