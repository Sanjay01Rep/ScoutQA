"""Load a template file (xlsx / csv / md / json) or the built-in default into a TemplateSpec."""

from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from scoutqa.config.models import TemplateConfig
from scoutqa.errors import ConfigError
from scoutqa.template.mapping import map_columns, normalise
from scoutqa.template.spec import CanonicalField, ColumnSpec, TemplateSpec

DEFAULT_COLUMNS = ["ID", "Module", "Scenario", "Preconditions", "Steps", "Expected Result", "Priority"]
DEFAULT_PRIORITIES = ["High", "Medium", "Low"]
_HEADER_SCAN_ROWS = 25


def load_template(cfg: TemplateConfig, base_dir: Path | None = None) -> TemplateSpec:
    if not cfg.path:
        return _finish(TemplateSpec(source="default", format="default",
                                    columns=_columns(DEFAULT_COLUMNS, cfg, {6: DEFAULT_PRIORITIES})), cfg)
    path = Path(cfg.path)
    if not path.is_absolute() and base_dir is not None:
        path = base_dir / path
    if not path.is_file():
        raise ConfigError(f"Template not found: {path}")
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        spec = _from_xlsx(path, cfg)
    elif suffix == ".csv":
        spec = _from_csv(path, cfg)
    elif suffix in (".md", ".markdown"):
        spec = _from_markdown(path, cfg)
    elif suffix == ".json":
        spec = _from_json(path, cfg)
    else:
        raise ConfigError(f"Unsupported template type {suffix!r}: use .xlsx, .csv, .md or .json")
    if not any(c.field not in (CanonicalField.BLANK, CanonicalField.CUSTOM) for c in spec.columns):
        raise ConfigError(f"No recognisable columns in {path}. Map them with template.columns in scoutqa.yaml.")
    return _finish(spec, cfg)


def _columns(headers: list[str], cfg: TemplateConfig, enums: dict[int, list[str]] | None = None) -> list[ColumnSpec]:
    return map_columns(headers, cfg.columns, cfg.defaults, enums)


def _finish(spec: TemplateSpec, cfg: TemplateConfig) -> TemplateSpec:
    has_step_rows = any(c.field is CanonicalField.STEP_NO for c in spec.columns)
    spec.layout = cfg.layout if cfg.layout != "auto" else ("step_per_row" if has_step_rows else "case_per_row")
    return spec


# ---------------------------------------------------------------- xlsx


def _from_xlsx(path: Path, cfg: TemplateConfig) -> TemplateSpec:
    wb = load_workbook(path)
    ws = wb[cfg.sheet] if cfg.sheet else wb.worksheets[0]
    if cfg.sheet and cfg.sheet not in wb.sheetnames:
        raise ConfigError(f"Sheet {cfg.sheet!r} not in {path.name} (sheets: {', '.join(wb.sheetnames)})")
    header_row = cfg.header_row or _detect_header_row(ws)
    headers = [str(c.value).strip() if c.value is not None else "" for c in ws[header_row]]
    while headers and not headers[-1]:
        headers.pop()
    enums = _validation_lists(ws, wb, header_row)
    return TemplateSpec(source=str(path), format="xlsx", columns=_columns(headers, cfg, enums), sheet=ws.title,
                        header_row=header_row)


def _detect_header_row(ws: Worksheet) -> int:
    """The first row (within the first rows) with the most recognised column names."""
    best, best_score = 1, -1
    for r in range(1, min(ws.max_row, _HEADER_SCAN_ROWS) + 1):
        values = [str(c.value) for c in ws[r] if isinstance(c.value, str) and c.value.strip()]
        if len(values) < 2:
            continue
        known = sum(1 for c in map_columns(values) if c.field not in (CanonicalField.CUSTOM,))
        if known > best_score:
            best, best_score = r, known
    return best


def _validation_lists(ws: Worksheet, wb: Any, header_row: int) -> dict[int, list[str]]:
    """Excel dropdowns (data validation lists) below the header -> allowed values per column index."""
    enums: dict[int, list[str]] = {}
    for dv in ws.data_validations.dataValidation:
        if dv.type != "list" or not dv.formula1:
            continue
        values = _list_values(dv.formula1, wb)
        if not values:
            continue
        for rng in dv.sqref.ranges:
            if rng.max_row is not None and rng.max_row <= header_row:
                continue
            for col in range(rng.min_col, rng.max_col + 1):
                enums.setdefault(col - 1, values)
    return enums


def _list_values(formula: str, wb: Any) -> list[str]:
    formula = formula.strip()
    if formula.startswith('"') and formula.endswith('"'):
        return [v.strip() for v in formula[1:-1].split(",") if v.strip()]
    ref = formula.lstrip("=")
    if "!" in ref:  # =Lists!$A$1:$A$3
        sheet, cells = ref.rsplit("!", 1)
        sheet = sheet.strip("'")
        if sheet in wb.sheetnames:
            rows = wb[sheet][cells.replace("$", "")]
            flat = [c for row in (rows if isinstance(rows, tuple) else ((rows,),)) for c in
                    (row if isinstance(row, tuple) else (row,))]
            return [str(c.value).strip() for c in flat if c.value is not None and str(c.value).strip()]
    return []


# ---------------------------------------------------------------- csv / md / json


def _from_csv(path: Path, cfg: TemplateConfig) -> TemplateSpec:
    with path.open(encoding="utf-8-sig", newline="") as fh:
        headers = next(csv.reader(fh), [])
    return TemplateSpec(source=str(path), format="csv", columns=_columns(headers, cfg))


def _from_markdown(path: Path, cfg: TemplateConfig) -> TemplateSpec:
    for line in path.read_text(encoding="utf-8").splitlines():
        stripped = line.strip()
        if stripped.startswith("|") and stripped.count("|") >= 3:
            headers = [h.strip() for h in stripped.strip("|").split("|")]
            return TemplateSpec(source=str(path), format="md", columns=_columns(headers, cfg))
    raise ConfigError(f"{path.name}: expected a Markdown table header like '| ID | Title | Steps |'")


def _from_json(path: Path, cfg: TemplateConfig) -> TemplateSpec:
    data = json.loads(path.read_text(encoding="utf-8"))
    enums: dict[int, list[str]] = {}
    if isinstance(data, dict) and isinstance(data.get("properties"), dict):  # JSON Schema
        headers = list(data["properties"])
        for i, name in enumerate(headers):
            prop = data["properties"][name]
            if isinstance(prop, dict) and isinstance(prop.get("enum"), list):
                enums[i] = [str(v) for v in prop["enum"]]
    elif isinstance(data, dict) and isinstance(data.get("columns"), list):
        headers = [str(c) for c in data["columns"]]
    elif isinstance(data, list):
        headers = [str(c) for c in data]
    else:
        raise ConfigError(f"{path.name}: expected a JSON Schema, a list of columns or {{\"columns\": [...]}}")
    return TemplateSpec(source=str(path), format="json", columns=_columns(headers, cfg, enums))


def describe(spec: TemplateSpec) -> list[tuple[str, str, str]]:
    """(column, mapped field, reason) rows for `scoutqa template`."""
    return [(c.name, c.field.value + (f" [{', '.join(c.enum)}]" if c.enum else "")
             + (f" = {c.default!r}" if c.default else ""), c.reason) for c in spec.columns]


__all__ = ["DEFAULT_COLUMNS", "describe", "load_template", "normalise"]
