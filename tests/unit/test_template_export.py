"""Template loading/mapping and all exporters, on hand-built cases (no browser)."""

from __future__ import annotations

import csv
import json
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from scoutqa.config.models import TemplateConfig
from scoutqa.errors import ConfigError
from scoutqa.export.excel import export_excel
from scoutqa.export.report import Report
from scoutqa.export.text_formats import export_csv, export_json, export_markdown
from scoutqa.generate.cases import CaseType, Priority, Source, Step, TestCase
from scoutqa.template.loader import load_template
from scoutqa.template.mapping import map_columns
from scoutqa.template.spec import CanonicalField as F
from scoutqa.template.values import priority_value, rows_for, scenario_of


def case(n: int = 1, title: str = "'Name' is required on 'Edit item' (/items/new)", **kw: object) -> TestCase:
    data: dict[str, object] = dict(
        key=f"k{n}", id=f"TC-ITEMS-{n:03d}", module="Items", title=title, type=CaseType.NEGATIVE,
        priority=Priority.HIGH, roles=["admin", "viewer"], preconditions=["User is signed in as any of: admin, viewer"],
        steps=[Step(action="Open the 'Edit item' page (/items/new)", data="/items/new"),
               Step(action="Leave 'Name' empty"),
               Step(action="Click 'Save'", expected="The form is not submitted")],
        expected_result="The form is not submitted and 'Name' shows: \"Please fill out this field.\"",
        source=Source(generator="R-FIELD-REQUIRED", state_id="s1", url_pattern="/items/new", refs=["e2"],
                      evidence=["'Name' is marked required (observed)"]),
    )
    data.update(kw)
    return TestCase.model_validate(data)


REPORT = Report(coverage=[["Items", "Edit item", "/items/new", "admin", 1, 0, 1, 4, 2, 10, 5, 3]],
                limitations=[["Scope of crawling", "Only reachable UI", "", ""]],
                trace=[["TC-ITEMS-001", "t", "rule", "R-FIELD-REQUIRED", "/items/new", "s1", "e2", "ev", "", "No"]],
                totals={"cases": 1, "high": 1, "needs_review": 0})


# ---------------------------------------------------------------- mapping

def test_default_template() -> None:
    spec = load_template(TemplateConfig())
    assert [(c.name, c.field) for c in spec.columns] == [
        ("ID", F.ID), ("Module", F.MODULE), ("Scenario", F.TITLE), ("Preconditions", F.PRECONDITIONS),
        ("Steps", F.STEPS), ("Expected Result", F.EXPECTED), ("Priority", F.PRIORITY)]
    assert spec.column(F.PRIORITY).enum == ["High", "Medium", "Low"]  # type: ignore[union-attr]
    assert spec.layout == "case_per_row"


@pytest.mark.parametrize(("header", "field"), [
    ("TC ID", F.ID), ("Test Case #", F.ID), ("Test Scenario", F.TITLE), ("Pre-Conditions", F.PRECONDITIONS),
    ("Test Steps", F.STEPS), ("Expected Result(s)", F.EXPECTED), ("Severity", F.PRIORITY), ("Test Data", F.TEST_DATA),
    ("Actual Result", F.BLANK), ("Status", F.BLANK), ("Executed By", F.BLANK), ("Defect ID", F.BLANK),
    ("Automation Candidate?", F.CUSTOM), ("Remarks", F.NOTES), ("User Role", F.ROLES), ("Screen", F.PAGE),
    ("Seq", F.STEP_NO), ("Sequence No", F.STEP_NO), ("Testcase Type", F.TYPE),
])
def test_synonyms(header: str, field: F) -> None:
    (column,) = map_columns([header])
    assert column.field is field, column.reason


def test_scenario_and_title_both_present() -> None:
    cols = {c.name: c.field for c in map_columns(["Scenario", "Test Case Title", "Steps"])}
    assert cols == {"Scenario": F.SCENARIO, "Test Case Title": F.TITLE, "Steps": F.STEPS}


def test_description_is_the_scenario_when_a_title_column_already_exists() -> None:
    # 'Description' alone (no title column) still falls back to the case title via
    # _resolve_title_vs_scenario; with both, it must not collide with 'Name'.
    cols = {c.name: c.field for c in map_columns(["Name", "Description"])}
    assert cols == {"Name": F.TITLE, "Description": F.SCENARIO}
    (only,) = map_columns(["Description"])
    assert only.field is F.TITLE


def test_step_per_row_detection() -> None:
    cols = {c.name: c.field for c in map_columns(["TC ID", "Title", "Step No", "Step Description", "Expected Result",
                                                    "Test Data"])}
    assert cols["Step No"] is F.STEP_NO and cols["Step Description"] is F.STEP_ACTION
    assert cols["Expected Result"] is F.STEP_EXPECTED and cols["Test Data"] is F.STEP_DATA


def test_fuzzy_match_respects_word_boundaries() -> None:
    # 'screen' must not fuzzy-match inside the unrelated word 'screenshot' (it's a coincidence of
    # spelling, not evidence this is a page/screen-name column); a plain header with no other steps-ish
    # column around legitimately still matches the whole word 'steps' later in the same header.
    (column,) = map_columns(["Mandate Screenshot (Type Yes for required steps)"])
    assert "screen" not in column.reason
    assert column.field is not F.PAGE


def test_step_per_row_second_steps_like_column_becomes_custom_not_a_collision() -> None:
    # A second, unrelated column that also happens to contain the whole word 'steps' (here, as part of
    # ordinary English rather than meaning "this is the steps column") must not silently share
    # STEP_ACTION with the real step-description column.
    cols = {c.name: c.field for c in map_columns(
        ["Seq", "Step Description", "Expected Result", "Mandate Screenshot (for required steps)"])}
    assert cols["Step Description"] is F.STEP_ACTION
    assert cols["Mandate Screenshot (for required steps)"] is F.CUSTOM


def test_overrides_and_defaults() -> None:
    cols = {c.name: c for c in map_columns(["Env", "Summary", "Owner"], overrides={"Owner": "blank"},
                                            defaults={"Env": "QA"})}
    assert cols["Env"].default == "QA" and cols["Owner"].field is F.BLANK
    with pytest.raises(ConfigError, match="unknown field"):
        map_columns(["X"], overrides={"X": "nonsense"})


def test_duplicate_mapping_becomes_custom() -> None:
    cols = [c.field for c in map_columns(["Title", "Summary"])]
    assert cols == [F.TITLE, F.CUSTOM]


# ---------------------------------------------------------------- loaders

def _xlsx_template(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Cases"
    ws["A1"] = "ACME test case template v2"
    ws.append([])
    ws.append(["TC ID", "Test Scenario", "Steps", "Expected Result", "Priority", "Status", "Automation Candidate"])
    ws.append(["TC-EX-1", "Example row", "1. do", "ok", "P2", "", "Yes"])
    for cell in ws[4]:
        cell.font = Font(italic=True, color="FF0000")
        cell.fill = PatternFill("solid", fgColor="FFF2CC")
    dv = DataValidation(type="list", formula1='"P1,P2,P3"')
    dv.add("E4:E10")
    ws.add_data_validation(dv)
    lists = wb.create_sheet("Lists")
    for i, v in enumerate(["Not run", "Pass", "Fail"], 1):
        lists.cell(row=i, column=1, value=v)
    status = DataValidation(type="list", formula1="=Lists!$A$1:$A$3")
    status.add("F4:F10")
    ws.add_data_validation(status)
    wb.save(path)


def test_xlsx_loader_detects_header_and_dropdowns(tmp_path: Path) -> None:
    path = tmp_path / "tpl.xlsx"
    _xlsx_template(path)
    spec = load_template(TemplateConfig(path=str(path)))
    assert spec.header_row == 3 and spec.sheet == "Cases"
    cols = {c.name: c for c in spec.columns}
    assert cols["Priority"].enum == ["P1", "P2", "P3"]
    assert cols["Status"].field is F.BLANK and cols["Status"].enum == ["Not run", "Pass", "Fail"]
    assert cols["Automation Candidate"].field is F.CUSTOM
    assert [c.name for c in spec.custom_columns] == ["Automation Candidate"]


def test_csv_md_json_loaders(tmp_path: Path) -> None:
    (tmp_path / "t.csv").write_text("ID,Title,Steps,Expected\n", encoding="utf-8-sig")
    (tmp_path / "t.md").write_text("# Template\n\n| ID | Title | Steps | Expected Result |\n|---|---|---|---|\n",
                                   encoding="utf-8")
    (tmp_path / "t.json").write_text(json.dumps({"properties": {"id": {}, "title": {}, "priority": {
        "enum": ["Critical", "Major", "Minor"]}}}), encoding="utf-8")
    for name in ("t.csv", "t.md"):
        spec = load_template(TemplateConfig(path=str(tmp_path / name)))
        assert [c.field for c in spec.columns][:3] == [F.ID, F.TITLE, F.STEPS]
    spec = load_template(TemplateConfig(path=str(tmp_path / "t.json")))
    assert spec.column(F.PRIORITY).enum == ["Critical", "Major", "Minor"]  # type: ignore[union-attr]


def test_loader_errors(tmp_path: Path) -> None:
    with pytest.raises(ConfigError, match="not found"):
        load_template(TemplateConfig(path=str(tmp_path / "missing.xlsx")))
    (tmp_path / "t.txt").write_text("x", encoding="utf-8")
    with pytest.raises(ConfigError, match="Unsupported"):
        load_template(TemplateConfig(path=str(tmp_path / "t.txt")))
    (tmp_path / "junk.csv").write_text("Foo,Bar\n", encoding="utf-8")
    with pytest.raises(ConfigError, match="No recognisable columns"):
        load_template(TemplateConfig(path=str(tmp_path / "junk.csv")))


# ---------------------------------------------------------------- values

@pytest.mark.parametrize(("priority", "allowed", "expected"), [
    (Priority.HIGH, [], "High"), (Priority.HIGH, ["P1", "P2", "P3"], "P1"), (Priority.LOW, ["P1", "P2", "P3"], "P3"),
    (Priority.MEDIUM, ["Critical", "Major", "Minor"], "Major"),  # no 'medium' synonym -> middle of the scale
    (Priority.HIGH, ["1 - Urgent", "2 - Normal", "3 - Low"], "1 - Urgent"),
    (Priority.MEDIUM, ["Blocker", "Critical", "Major", "Minor", "Trivial"], "Major"),
])
def test_priority_mapping(priority: Priority, allowed: list[str], expected: str) -> None:
    assert priority_value(priority, allowed) == expected


def test_case_per_row_values() -> None:
    spec = load_template(TemplateConfig())
    (row,) = rows_for(case(), spec)
    values = dict(zip([c.name for c in spec.columns], row, strict=True))
    assert values["ID"] == "TC-ITEMS-001" and values["Priority"] == "High"
    # step data already visible in the action text is not repeated in brackets
    assert values["Steps"] == "1. Open the 'Edit item' page (/items/new)\n2. Leave 'Name' empty\n3. Click 'Save'"
    assert values["Preconditions"] == "User is signed in as any of: admin, viewer"
    assert scenario_of(case()) == "Form input validation — /items/new"


def test_step_per_row_values() -> None:
    spec = load_template(TemplateConfig())
    spec.columns = map_columns(["ID", "Title", "Step #", "Action", "Expected Result", "Status"])
    spec.layout = "step_per_row"
    rows = rows_for(case(), spec)
    assert len(rows) == 3
    assert [r[2] for r in rows] == ["1", "2", "3"]
    assert rows[0][1] and rows[1][1] == ""  # case columns only on the first row
    assert all(r[0] == "TC-ITEMS-001" for r in rows)  # ID repeated for filtering
    assert rows[2][4] == "The form is not submitted" and rows[2][5] == ""


# ---------------------------------------------------------------- exporters

def test_excel_default_template(tmp_path: Path) -> None:
    out = tmp_path / "out.xlsx"
    rows = export_excel([case(1), case(2, title="=HYPERLINK(\"http://evil\")")], load_template(TemplateConfig()),
                        REPORT, out)
    assert rows == 2
    wb = load_workbook(out)
    assert wb.sheetnames == ["Test Cases", "Coverage", "Limitations", "Trace"]
    ws = wb["Test Cases"]
    assert [c.value for c in ws[1]] == ["ID", "Module", "Scenario", "Preconditions", "Steps", "Expected Result",
                                        "Priority"]
    assert ws["A2"].value == "TC-ITEMS-001" and ws["C3"].data_type == "s"  # formula stays text
    assert ws["C3"].value.startswith("=HYPERLINK")
    assert ws.freeze_panes == "B2" and ws.auto_filter.ref == "A1:G3"
    assert any("G2" in str(dv.sqref) for dv in ws.data_validations.dataValidation)
    assert wb["Coverage"]["C2"].value == "/items/new"


def test_excel_preserves_the_users_template(tmp_path: Path) -> None:
    tpl, out = tmp_path / "tpl.xlsx", tmp_path / "out.xlsx"
    _xlsx_template(tpl)
    spec = load_template(TemplateConfig(path=str(tpl)))
    export_excel([case(n) for n in range(1, 13)], spec, REPORT, out)
    wb = load_workbook(out)
    assert wb.sheetnames == ["Cases", "Lists", "Coverage", "Limitations", "Trace"]
    ws = wb["Cases"]
    assert ws["A1"].value == "ACME test case template v2"  # title rows kept
    assert ws["A4"].value == "TC-ITEMS-001" and ws["A15"].value == "TC-ITEMS-012"  # example row replaced
    assert ws["E4"].value == "P1"  # High -> the template's own scale
    assert ws["F4"].value in (None, "")  # execution column left blank
    assert ws["B5"].font.italic and ws["B5"].fill.fgColor.rgb.endswith("FFF2CC")  # example row style reused
    ranges = " ".join(str(dv.sqref) for dv in ws.data_validations.dataValidation)
    assert "E4:E15" in ranges and "F4:F15" in ranges  # dropdowns stretched to every written row


def test_csv_export_is_excel_friendly_and_injection_safe(tmp_path: Path) -> None:
    out = tmp_path / "out.csv"
    export_csv([case(1, title="=cmd|' /C calc'!A0")], load_template(TemplateConfig()), out)
    raw = out.read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")
    rows = list(csv.reader(out.open(encoding="utf-8-sig")))
    assert rows[0][0] == "ID" and rows[1][2].startswith("'=")
    assert "\n" in rows[1][4]  # multi-line steps survive in one cell


def test_markdown_export(tmp_path: Path) -> None:
    out = tmp_path / "out.md"
    export_markdown([case()], load_template(TemplateConfig()), REPORT, out, "demo")
    text = out.read_text(encoding="utf-8")
    assert "## Items" in text and "### TC-ITEMS-001" in text
    assert "| 3 | Click 'Save' |   | The form is not submitted |" in text
    assert "## Limitations" in text


def test_markdown_template_gives_one_table(tmp_path: Path) -> None:
    tpl = tmp_path / "t.md"
    tpl.write_text("| ID | Title | Steps |\n|---|---|---|\n", encoding="utf-8")
    out = tmp_path / "out.md"
    export_markdown([case()], load_template(TemplateConfig(path=str(tpl))), REPORT, out, "demo")
    assert "| TC-ITEMS-001 | 'Name' is required on 'Edit item' (/items/new) | 1. Open" in out.read_text(
        encoding="utf-8")


def test_json_export_is_lossless(tmp_path: Path) -> None:
    out = tmp_path / "out.json"
    export_json([case()], load_template(TemplateConfig()), REPORT, out)
    data = json.loads(out.read_text(encoding="utf-8"))
    assert TestCase.model_validate(data["cases"][0]) == case()
    assert data["rows"][0]["ID"] == "TC-ITEMS-001"
    assert data["limitations"][0]["Area"] == "Scope of crawling"
