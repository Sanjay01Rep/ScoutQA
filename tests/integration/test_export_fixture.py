"""Milestone 4 end-to-end: crawl -> generate (rules) -> export in every format. Zero tokens throughout."""

from __future__ import annotations

import csv
import json
from collections.abc import Callable
from pathlib import Path

import pytest
from openpyxl import load_workbook

from scoutqa import pipeline
from scoutqa.config.models import ProjectConfig
from scoutqa.errors import ScoutQAError
from tests.fixture_app.server import PASSWORD, FixtureServer

pytestmark = pytest.mark.browser

ConfigFactory = Callable[..., ProjectConfig]


@pytest.fixture
async def generated(app_server: FixtureServer, make_config: ConfigFactory, creds: tuple[str, str]) -> ProjectConfig:
    cfg = make_config()
    await pipeline.crawl(cfg)
    pipeline.generate(cfg)
    return cfg


async def test_end_to_end_excel(generated: ProjectConfig, tmp_path: Path) -> None:
    out = tmp_path / "cases.xlsx"
    report = pipeline.export(generated, fmt="xlsx", output=out)
    assert report.cases > 40 and report.rows == report.cases and report.template == "default"
    wb = load_workbook(out)
    assert wb.sheetnames == ["Test Cases", "Coverage", "Limitations", "Trace"]
    ws = wb["Test Cases"]
    ids = [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)]
    assert len(ids) == report.cases and len(set(ids)) == len(ids)
    titles = {ws.cell(row=r, column=3).value for r in range(2, ws.max_row + 1)}
    assert "'Name' is required on 'Edit item' (/items/new)" in titles
    assert {ws.cell(row=r, column=7).value for r in range(2, ws.max_row + 1)} <= {"High", "Medium", "Low"}

    coverage = {row[2]: row for row in wb["Coverage"].iter_rows(min_row=2, values_only=True)}
    items_new = coverage["/items/new"]
    assert items_new[7] == 4 and items_new[8] == 4  # 4 fields, all 4 covered by at least one case
    limitations = [row[0] for row in wb["Limitations"].iter_rows(min_row=2, values_only=True)]
    assert "Scope of crawling" in limitations and "Forms never submitted" in limitations
    assert any(str(a).startswith("Requests blocked") for a in limitations)  # the /api/track beacon
    assert any(str(a).startswith("Unsafe links not followed") for a in limitations)
    trace = list(wb["Trace"].iter_rows(min_row=2, values_only=True))
    assert len(trace) == report.cases and all(row[3].startswith("R-") for row in trace)


async def test_end_to_end_other_formats(generated: ProjectConfig, tmp_path: Path) -> None:
    for fmt in ("csv", "md", "json"):
        out = tmp_path / f"cases.{fmt}"
        report = pipeline.export(generated, fmt=fmt, output=out)  # type: ignore[arg-type]
        text = out.read_text(encoding="utf-8-sig")
        assert PASSWORD not in text
        assert "TC-ITEMS-" in text, fmt
        if fmt == "json":
            assert len(json.loads(text)["cases"]) == report.cases


async def test_export_uses_configured_template(generated: ProjectConfig, tmp_path: Path) -> None:
    tpl = tmp_path / "team.csv"
    tpl.write_text("TC #,Test Scenario,Pre-conditions,Test Steps,Expected Result,Severity,Status,Actual Result\n",
                   encoding="utf-8")
    cfg = generated.model_copy(update={"template": generated.template.model_copy(update={"path": str(tpl)})})
    out = tmp_path / "team-out.csv"
    report = pipeline.export(cfg, fmt="csv", output=out)
    with out.open(encoding="utf-8-sig", newline="") as fh:
        header, *rows = list(csv.reader(fh))
    assert header == ["TC #", "Test Scenario", "Pre-conditions", "Test Steps", "Expected Result", "Severity",
                      "Status", "Actual Result"]
    assert len(rows) == report.cases and all(r[0].startswith("TC-") for r in rows)
    assert all(r[6] == "" and r[7] == "" for r in rows)  # execution-time columns stay blank
    assert all(r[5] in ("High", "Medium", "Low") for r in rows)
    assert "User is signed in" in {r[2] for r in rows}  # single unnamed profile: no 'default' role in text
    assert report.custom_columns == []


def test_export_before_generate_fails(make_config: ConfigFactory) -> None:
    with pytest.raises(ScoutQAError, match="generate"):
        pipeline.export(make_config())
